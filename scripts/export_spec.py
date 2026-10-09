#!/usr/bin/env python3
"""Generate Solaris specification PDF and DOCX from one tracked Markdown source.

Build on GitHub Windows runner (Arial is installed). No font files distributed.
"""
from __future__ import annotations

import re
import sys
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "SOLARIS-NEON-2.2.10-SPEC.md"
OUT = ROOT / "artifacts" / "spec"
OUT.mkdir(parents=True, exist_ok=True)
BASENAME = "SOLARIS-NEON-2.2.10-TECHNICAL-SPEC"
GOLD = (183, 117, 19)
INK = (30, 29, 36)

def font_candidates():
    win = Path("C:/Windows/Fonts")
    for regular, bold in [
        (win / "arial.ttf", win / "arialbd.ttf"),
        (win / "segoeui.ttf", win / "segoeuib.ttf"),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
    ]:
        if regular.exists() and bold.exists():
            return regular, bold
    raise RuntimeError("A Unicode-capable font is required for Russian PDF output.")

def parse_source():
    blocks = []
    para_lines = []
    def flush():
        nonlocal para_lines
        if para_lines:
            blocks.append(("paragraph", " ".join(para_lines)))
            para_lines = []
    for line in SOURCE.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        header = re.match(r"^(#{1,4})\s+(.+)", stripped)
        if header:
            flush()
            blocks.append(("h" + str(len(header.group(1))), header.group(2)))
        elif stripped.startswith("> "):
            flush()
            blocks.append(("quote", stripped[2:]))
        elif re.match(r"^[-*]\s+", stripped):
            flush()
            blocks.append(("bullet", re.sub(r"^[-*]\s+", "", stripped)))
        elif re.match(r"^\d+[.]\s+", stripped):
            flush()
            blocks.append(("number", stripped))
        elif stripped == "---":
            flush()
            blocks.append(("rule", ""))
        else:
            para_lines.append(stripped)
    flush()
    return blocks

def segments(markup: str):
    # preserve bold and inline code, otherwise pass through verbatim.
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", markup)
    for token in parts:
        if token.startswith("**") and token.endswith("**"):
            yield "bold", token[2:-2]
        elif token.startswith("`") and token.endswith("`"):
            yield "code", token[1:-1]
        else:
            yield "text", token

def pdf_text(markup: str) -> str:
    result = []
    for kind, value in segments(markup):
        clean = escape(value)
        if kind == "bold":
            result.append("<b>" + clean + "</b>")
        elif kind == "code":
            result.append('<font color="#85520E">' + clean + "</font>")
        else:
            result.append(clean)
    return "".join(result)

def build_pdf(blocks):
    normal, bold = font_candidates()
    pdfmetrics.registerFont(TTFont("Solaris", str(normal)))
    pdfmetrics.registerFont(TTFont("Solaris-Bold", str(bold)))
    pdfmetrics.registerFontFamily("Solaris", normal="Solaris", bold="Solaris-Bold")

    styles = {
        "paragraph": ParagraphStyle(
            "body", fontName="Solaris", fontSize=9.5, leading=14.5,
            spaceAfter=6, textColor=colors.HexColor("#26222A")),
        "h1": ParagraphStyle(
            "title", fontName="Solaris-Bold", fontSize=18, leading=24,
            spaceAfter=16, textColor=colors.HexColor("#9C650E")),
        "h2": ParagraphStyle(
            "section", fontName="Solaris-Bold", fontSize=13.5, leading=19,
            spaceBefore=13, spaceAfter=7, keepWithNext=True,
            textColor=colors.HexColor("#996113")),
        "h3": ParagraphStyle(
            "subsection", fontName="Solaris-Bold", fontSize=11.5, leading=16,
            spaceBefore=11, spaceAfter=6, keepWithNext=True,
            textColor=colors.HexColor("#504052")),
        "h4": ParagraphStyle(
            "smallheader", fontName="Solaris-Bold", fontSize=10, leading=14,
            spaceBefore=8, spaceAfter=4, keepWithNext=True),
        "bullet": ParagraphStyle(
            "bullet", fontName="Solaris", fontSize=9.5, leading=14.5,
            leftIndent=18, firstLineIndent=-9, spaceAfter=4),
        "number": ParagraphStyle(
            "number", fontName="Solaris", fontSize=9.5, leading=14.5,
            leftIndent=18, firstLineIndent=-9, spaceAfter=4),
        "quote": ParagraphStyle(
            "quote", fontName="Solaris", fontSize=9.2, leading=14,
            leftIndent=12, borderColor=colors.HexColor("#CE9B49"),
            borderWidth=1, borderPadding=8, spaceAfter=10,
            textColor=colors.HexColor("#795317"))
    }
    doc = SimpleDocTemplate(
        str(OUT / (BASENAME + ".pdf")), pagesize=A4,
        leftMargin=48, rightMargin=48, topMargin=64, bottomMargin=55,
        title="SOLARIS NEON 2.2.10 — Техническое задание",
        author="SOLARIS LAUNCHER Project")
    flow = []
    for kind, body in blocks:
        if kind == "rule":
            flow.append(Spacer(1, 8))
            continue
        text = pdf_text(body)
        if kind == "bullet":
            text = "• " + text
        flow.append(Paragraph(text, styles[kind]))

    def page_decor(canvas, doc):
        canvas.saveState()
        w, h = A4
        canvas.setStrokeColor(colors.HexColor("#E3C080"))
        canvas.setLineWidth(0.8)
        canvas.line(48, h - 44, w - 48, h - 44)
        canvas.setFont("Solaris", 8)
        canvas.setFillColor(colors.HexColor("#837482"))
        canvas.drawString(48, 35, "SOLARIS NEON 2.2.10  |  TECHNICAL SPEC")
        canvas.drawRightString(w - 48, 35, f"{doc.page}")
        canvas.restoreState()

    doc.build(flow, onFirstPage=page_decor, onLaterPages=page_decor)

def append_docx_inline(paragraph, content):
    for kind, text in segments(content):
        run = paragraph.add_run(text)
        if kind == "bold": run.bold = True
        if kind == "code":
            run.font.color.rgb = RGBColor(133, 82, 14)
            run.font.name = "Consolas"

def build_docx(blocks):
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(*INK)
    normal.paragraph_format.space_after = Pt(5)
    for level, size in [(1, 19), (2, 14), (3, 12), (4, 10.5)]:
        style = document.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*GOLD) if level <= 2 else RGBColor(67, 52, 74)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("☀  SOLARIS  /  NEON 2.2.10")
    run.font.name = "Arial"
    run.font.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(*GOLD)

    for kind, body in blocks:
        if kind.startswith("h"):
            level = int(kind[1])
            p = document.add_paragraph(style=f"Heading {level}")
        elif kind == "bullet":
            p = document.add_paragraph(style="List Bullet")
        elif kind == "number":
            p = document.add_paragraph(style="List Number")
            body = re.sub(r"^\d+[.]\s+", "", body)
        else:
            p = document.add_paragraph()
            if kind == "quote":
                p.paragraph_format.left_indent = Cm(0.5)
                p.style = document.styles["Quote"]
        if kind == "rule":
            p.add_run("—" * 24)
            continue
        append_docx_inline(p, body)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("SOLARIS LAUNCHER  •  10 из 10 продуктовых пунктов утверждены")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(110, 103, 115)
    document.save(OUT / (BASENAME + ".docx"))

def main():
    blocks = parse_source()
    assert len(blocks) > 60, "Unexpectedly incomplete source specification"
    build_docx(blocks)
    build_pdf(blocks)
    for path in sorted(OUT.glob(BASENAME + ".*")):
        print(f"CREATED {path.name}: {path.stat().st_size:,} bytes")

if __name__ == "__main__":
    main()
